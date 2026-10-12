"""Versión en diapositivas de un video: un paso narrado por diapositiva.

Desde la raíz del repositorio:

    uv run python -m herramientas_video.diapositivas entrega-1/video/video_entrega_1.py

Recorre la escena VideoCompleto sin animar, guarda una imagen al final de cada paso
narrado y arma dos archivos junto al video:

    diapositivas.pptx   cada imagen a pantalla completa; la narración de ese paso
                        queda en las notas del orador
    diapositivas.pdf    las mismas imágenes, sin notas

Son los que se suben al repositorio con cada entrega. Hay que rehacerlos cuando el
video cambia. No genera voz ni vuelve a renderizar el video.
"""

import argparse
from pathlib import Path

import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Inches

from .creditos import INTEGRANTES
from .guion import recorrer
from .voz import RAIZ

AUTORES = ", ".join(INTEGRANTES)


def diapositivas(carpeta, pasos):
    """(imagen, nota) por diapositiva; las capturas repetidas se funden en una."""
    resultado, anterior = [], None
    for numero, paso in enumerate(pasos, start=1):
        nota = paso["nota"]
        if nota and paso["quien"]:
            nota = f"{paso['quien']}: {nota}"
        ruta = carpeta / f"{numero:03d}.png"
        pixeles = np.asarray(Image.open(ruta).convert("RGB"))
        if anterior is not None and np.array_equal(pixeles, anterior):
            if nota:
                imagen, previa = resultado[-1]
                resultado[-1] = (imagen, f"{previa} {nota}".strip())
            continue
        resultado.append((ruta, nota))
        anterior = pixeles
    return resultado


def armar_pptx(titulo, laminas, destino):
    presentacion = Presentation()
    presentacion.slide_width = Inches(13.333)
    presentacion.slide_height = Inches(7.5)
    presentacion.core_properties.title = titulo
    presentacion.core_properties.author = AUTORES
    presentacion.core_properties.language = "es-CO"
    en_blanco = presentacion.slide_layouts[6]
    for numero, (imagen, nota) in enumerate(laminas, start=1):
        lamina = presentacion.slides.add_slide(en_blanco)
        foto = lamina.shapes.add_picture(
            str(imagen), Emu(0), Emu(0),
            width=presentacion.slide_width, height=presentacion.slide_height,
        )
        # texto alternativo: lo que se dice en ese paso
        descripcion = nota or f"{titulo}, paso {numero}"
        foto._element.xpath("./p:nvPicPr/p:cNvPr")[0].set("descr", descripcion)
        if nota:
            lamina.notes_slide.notes_text_frame.text = nota
    presentacion.save(destino)


def armar_pdf(titulo, laminas, destino):
    paginas = [Image.open(imagen).convert("RGB") for imagen, _ in laminas]
    # 1920 px a 144 ppp = 13,33 pulgadas: una página 16:9 del tamaño de la diapositiva
    paginas[0].save(
        destino, save_all=True, append_images=paginas[1:], resolution=144.0,
        title=titulo, author=AUTORES,
    )


def main():
    opciones = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    opciones.add_argument("video", help="archivo .py del video, con la escena VideoCompleto")
    opciones.add_argument("--escena", default="VideoCompleto")
    args = opciones.parse_args()

    archivo = Path(args.video).resolve()
    carpeta, recorrido = recorrer(archivo, args.escena, calidad="-qh")
    titulo = recorrido["titulo"] or archivo.stem
    laminas = diapositivas(carpeta, recorrido["pasos"])
    armar_pptx(titulo, laminas, archivo.parent / "diapositivas.pptx")
    armar_pdf(titulo, laminas, archivo.parent / "diapositivas.pdf")
    print(f"{len(laminas)} diapositivas -> {archivo.parent.relative_to(RAIZ)}")


if __name__ == "__main__":
    main()
